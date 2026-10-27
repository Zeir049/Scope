"""docx 업로드 파싱 + 미지원 확장자 테스트."""
from __future__ import annotations

import io
import uuid

import docx
from httpx import AsyncClient

from app.core.db import SessionLocal
from app.repository.chapter_repo import ChapterRepository


async def test_docx_upload_extracts_text(client: AsyncClient, auth_headers: dict, project_id: str):
    buf = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("도끼 이야기 첫 문단")
    document.add_paragraph("두 번째 문단")
    document.save(buf)

    r = await client.post(
        f"/projects/{project_id}/chapters",
        data={"chapter_no": "1", "title": "1화"},
        files={
            "file": (
                "ch1.docx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=auth_headers,
    )
    assert r.status_code == 201

    # 본문이 텍스트로 추출되어 저장됐는지 확인
    async with SessionLocal() as s:
        chapter = await ChapterRepository(s).get(uuid.UUID(project_id), 1)
        assert "도끼 이야기 첫 문단" in (chapter.raw_text or "")
        assert "두 번째 문단" in (chapter.raw_text or "")


async def test_unsupported_extension_400(client: AsyncClient, auth_headers: dict, project_id: str):
    r = await client.post(
        f"/projects/{project_id}/chapters",
        data={"chapter_no": "2"},
        files={"file": ("ch.exe", b"\x00binary", "application/octet-stream")},
        headers=auth_headers,
    )
    assert r.status_code == 400
